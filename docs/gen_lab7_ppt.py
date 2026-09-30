# -*- coding: utf-8 -*-
"""
实验七答辩 PPT 生成脚本
选题：项目九·微信小程序应用 —— 校园报修助手
作者：李春雨（学号 202305567128，计科 2 班）
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------- 全局样式 ----------
PRIMARY = RGBColor(0x67, 0x50, 0xA4)      # 主题紫（与小程序导航栏一致）
ACCENT = RGBColor(0xFF, 0x7A, 0x45)       # 强调橙
DARK = RGBColor(0x21, 0x21, 0x21)
GRAY = RGBColor(0x60, 0x60, 0x60)
LIGHT_BG = RGBColor(0xF6, 0xF6, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT_TITLE = "微软雅黑"
FONT_BODY = "微软雅黑"

prs = Presentation()
# 16:9
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

BLANK = prs.slide_layouts[6]

# ---------- 工具函数 ----------
def add_slide():
    return prs.slides.add_slide(BLANK)

def fill_shape(slide, color):
    """整页背景"""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    bg.shadow.inherit = False
    return bg

def add_rect(slide, x, y, w, h, color):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background()
    r.shadow.inherit = False
    return r

def add_text(slide, x, y, w, h, text, size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT,
             font=FONT_BODY, anchor=MSO_ANCHOR.TOP, line_spacing=1.25):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = font
    return tb

def _pt(v):
    """size 可能传裸数字或 Pt/Emu，统一转 Pt"""
    if hasattr(v, "emu"):
        return v
    return Pt(int(v))

def add_header(slide, title, page_no=None, total=10):
    """页眉：左侧紫色色块 + 标题"""
    add_rect(slide, 0, 0, Inches(0.35), Inches(7.5), PRIMARY)
    add_rect(slide, 0, 0, SW, Inches(1.1), WHITE)
    add_rect(slide, Inches(0.35), Inches(0.4), Inches(0.18), Inches(0.5), PRIMARY)
    add_text(slide, Inches(0.7), Inches(0.35), Inches(10), Inches(0.6),
             title, size=26, color=DARK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(11.0), Inches(0.4), Inches(2.0), Inches(0.5),
             f"{page_no or 0} / {total}", size=12, color=GRAY, align=PP_ALIGN.RIGHT,
             anchor=MSO_ANCHOR.MIDDLE)
    # 分隔线
    add_rect(slide, Inches(0.35), Inches(1.1), Inches(12.98), Emu(12700), RGBColor(0xE0, 0xE0, 0xE0))

def add_bullet(slide, x, y, w, items, size=18, color=DARK, gap=Pt(8)):
    tb = slide.shapes.add_textbox(x, y, w, Inches(5))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.4
        p.space_after = gap
        r1 = p.add_run(); r1.text = "▸ "
        r1.font.size = Pt(size); r1.font.color.rgb = ACCENT; r1.font.bold = True; r1.font.name = FONT_BODY
        r2 = p.add_run(); r2.text = it
        r2.font.size = Pt(size); r2.font.color.rgb = color; r2.font.name = FONT_BODY
    return tb

# =========================================================
# 第 1 页：封面
# =========================================================
s = add_slide()
fill_shape(s, PRIMARY)
# 装饰圆
for cx, cy, d in [(11.8, 1.0, 2.5), (12.5, 6.0, 1.8), (0.5, 6.5, 1.2)]:
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(cy), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0x7B, 0x5F, 0xB0)
    c.line.fill.background(); c.shadow.inherit = False

add_text(s, Inches(1.0), Inches(2.2), Inches(11), Inches(1.0),
         "校园报修助手", size=54, color=WHITE, bold=True)
add_text(s, Inches(1.0), Inches(3.3), Inches(11), Inches(0.6),
         "微信小程序 · 移动应用综合项目", size=28, color=RGBColor(0xE0, 0xCB, 0xF0))
add_rect(s, Inches(1.0), Inches(4.1), Inches(1.5), Emu(38100), ACCENT)
add_text(s, Inches(1.0), Inches(4.35), Inches(11), Inches(0.5),
         "《移动应用开发》实验七 · 综合项目答辩", size=20, color=WHITE)
add_text(s, Inches(1.0), Inches(5.0), Inches(11), Inches(0.5),
         "项目九·微信小程序应用", size=18, color=RGBColor(0xCF, 0xBC, 0xEC))
add_text(s, Inches(1.0), Inches(6.0), Inches(11), Inches(0.8),
         "李春雨 · 学号 202305567128 · 计科 2 班\nGitHub: lcyspring/Android (lcy 分支)",
         size=16, color=WHITE, line_spacing=1.5)

# =========================================================
# 第 2 页：选题与背景
# =========================================================
s = add_slide(); fill_shape(s, WHITE)
add_header(s, "一、选题与背景", 2)
add_text(s, Inches(0.7), Inches(1.4), Inches(12), Inches(0.6),
         "为什么做校园报修小程序？", size=22, color=PRIMARY, bold=True)

add_bullet(s, Inches(0.7), Inches(2.2), Inches(6.0), [
    "校园宿舍/教室/公共区域报修场景高频、刚需",
    "传统报修：到宿管处登记、电话沟通，记录易丢失",
    "维修进度不透明，无法及时反馈",
    "管理员缺少集中处理入口，统计不便",
], size=18)

# 右侧：解决方案卡片
card = add_rect(s, Inches(7.2), Inches(2.2), Inches(5.5), Inches(4.2), LIGHT_BG)
add_rect(s, Inches(7.2), Inches(2.2), Inches(5.5), Inches(0.6), PRIMARY)
add_text(s, Inches(7.4), Inches(2.25), Inches(5), Inches(0.5),
        "小程序方案优势", size=18, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
add_bullet(s, Inches(7.5), Inches(3.0), Inches(5.0), [
    "扫码即用，无需下载安装",
    "提交即存，记录不丢失",
    "处理完成自动订阅消息通知",
    "管理员集中处理入口",
    "云开发零后端运维",
], size=16)

add_text(s, Inches(0.7), Inches(6.6), Inches(12), Inches(0.5),
         "→ 选题对应任务书「项目九·微信小程序应用」，8 项基本功能全覆盖",
         size=15, color=ACCENT, bold=True)

# =========================================================
# 第 3 页：需求分析
# =========================================================
s = add_slide(); fill_shape(s, WHITE)
add_header(s, "二、需求分析", 3)

add_text(s, Inches(0.7), Inches(1.4), Inches(12), Inches(0.6),
         "核心用户与场景", size=22, color=PRIMARY, bold=True)

# 用户角色
roles = [("报修人", "学生/教师", PRIMARY), ("管理员", "宿管/后勤", ACCENT), ("访客", "浏览查看", GRAY)]
for i, (n, d, c) in enumerate(roles):
    x = Inches(0.7 + i * 4.3)
    add_rect(s, x, Inches(2.2), Inches(3.8), Inches(1.5), c)
    add_text(s, x, Inches(2.3), Inches(3.8), Inches(0.6), n, size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x, Inches(2.9), Inches(3.8), Inches(0.6), d, size=16, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# 核心功能清单
add_text(s, Inches(0.7), Inches(4.1), Inches(12), Inches(0.5),
         "核心功能（任务书 8 项全覆盖）", size=20, color=PRIMARY, bold=True)

funcs = [
    "① 多页面导航（4 页）", "② 用户输入+校验", "③ 列表展示+分页", "④ 数据增删改查",
    "⑤ 网络请求（云API）", "⑥ 本地数据保存", "⑦ 订阅消息通知", "⑧ 完整业务流程",
]
for i, t in enumerate(funcs):
    row, col = i // 4, i % 4
    x, y = Inches(0.7 + col * 3.15), Inches(4.7 + row * 0.95)
    add_rect(s, x, y, Inches(2.95), Inches(0.75), LIGHT_BG)
    add_rect(s, x, y, Emu(50800), Inches(0.75), PRIMARY if i % 2 == 0 else ACCENT)
    add_text(s, x + Inches(0.15), y, Inches(2.8), Inches(0.75), t,
             size=15, color=DARK, anchor=MSO_ANCHOR.MIDDLE)

# =========================================================
# 第 4 页：系统架构
# =========================================================
s = add_slide(); fill_shape(s, WHITE)
add_header(s, "三、系统架构", 4)

add_text(s, Inches(0.7), Inches(1.4), Inches(12), Inches(0.5),
        "分层架构：UI → Repository → 云/本地双模", size=20, color=PRIMARY, bold=True)

# 三层架构图
layers = [
    ("UI 层（Page）", "index 列表 / publish 报修 / detail 详情 / admin 管理", PRIMARY, Inches(1.6)),
    ("Repository 层（utils/api.js）", "统一封装所有数据访问，页面不直接调 wx.cloud", ACCENT, Inches(3.1)),
    ("数据源（双模）", "云开发（数据库/云函数/云存储）+ 本地兜底（wx.storage）", RGBColor(0x55, 0x7A, 0xB0), Inches(4.6)),
]
for name, desc, c, y in layers:
    add_rect(s, Inches(1.0), y, Inches(11.3), Inches(1.25), c)
    add_text(s, Inches(1.3), y + Inches(0.12), Inches(4), Inches(0.5),
             name, size=20, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(1.3), y + Inches(0.65), Inches(10.8), Inches(0.5),
             desc, size=15, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    if y < Inches(4.6):
        # 向下箭头
        arr = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.4), y + Inches(1.3), Inches(0.5), Inches(0.2))
        arr.fill.solid(); arr.fill.fore_color.rgb = GRAY; arr.line.fill.background()

# 实验七完善点
add_rect(s, Inches(1.0), Inches(6.0), Inches(11.3), Inches(1.2), LIGHT_BG)
add_rect(s, Inches(1.0), Inches(6.0), Emu(50800), Inches(1.2), ACCENT)
add_text(s, Inches(1.25), Inches(6.05), Inches(11), Inches(0.4),
         "实验七新增·双模 Repository", size=16, color=ACCENT, bold=True)
add_text(s, Inches(1.25), Inches(6.45), Inches(11), Inches(0.7),
         "USE_MOCK_FALLBACK 开关：无云环境自动回退 wx.storage，mock 种子 3 条 + 「我的报修」本地缓存，\n保证无云配额也能演示完整业务流程，覆盖任务书「网络请求」「本地数据保存」双要求。",
         size=14, color=DARK, line_spacing=1.3)

# =========================================================
# 第 5 页：页面与流程
# =========================================================
s = add_slide(); fill_shape(s, WHITE)
add_header(s, "四、页面与业务流程", 5)

add_text(s, Inches(0.7), Inches(1.4), Inches(12), Inches(0.5),
         "4 个核心页面 · 完整业务闭环", size=20, color=PRIMARY, bold=True)

# 页面卡片
pages = [
    ("index", "报修列表", "状态/楼栋筛选\n分页加载\n点击进详情", PRIMARY),
    ("publish", "提交报修", "表单录入\n输入校验\n选图+订阅消息", ACCENT),
    ("detail", "报修详情", "查看单条\n预览图片\n拨号/复制", RGBColor(0x4C, 0xAF, 0x50)),
    ("admin", "管理后台", "处理/删除\n权限判断\n处理通知", RGBColor(0x55, 0x7A, 0xB0)),
]
for i, (n, t, d, c) in enumerate(pages):
    x = Inches(0.7 + i * 3.15)
    add_rect(s, x, Inches(2.2), Inches(2.95), Inches(2.8), LIGHT_BG)
    add_rect(s, x, Inches(2.2), Inches(2.95), Inches(0.55), c)
    add_text(s, x, Inches(2.25), Inches(2.95), Inches(0.5), t,
             size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x, Inches(2.85), Inches(2.95), Inches(0.4), n,
             size=13, color=GRAY, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.15), Inches(3.3), Inches(2.65), Inches(1.6), d,
             size=14, color=DARK, line_spacing=1.4)

# 流程箭头
flow = "提交报修 → 写库/存本地 → 通知管理员 → 管理员处理 → 通知报修人 → 列表更新"
add_rect(s, Inches(0.7), Inches(5.4), Inches(11.95), Inches(0.9), PRIMARY)
add_text(s, Inches(0.7), Inches(5.4), Inches(11.95), Inches(0.9), flow,
         size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

add_text(s, Inches(0.7), Inches(6.5), Inches(12), Inches(0.6),
         "状态机：未处理 ⇄ 已处理（管理员处理触发切换，订阅消息通知报修人）",
         size=15, color=GRAY, align=PP_ALIGN.CENTER)

# =========================================================
# 第 6 页：关键技术实现
# =========================================================
s = add_slide(); fill_shape(s, WHITE)
add_header(s, "五、关键技术实现", 6)

add_text(s, Inches(0.7), Inches(1.4), Inches(12), Inches(0.5),
         "6 个核心代码片段", size=20, color=PRIMARY, bold=True)

snippets = [
    ("① 统一 Repository（utils/api.js）", "所有数据访问集中封装，页面不直接调 wx.cloud，方便维护与切换数据源", "「UI-ViewModel-Repository」架构"),
    ("② 云+本地双模回退（实验七新增）", "wx.cloud 不可用或调用失败时自动回退 wx.storage，mock 种子保证无云环境可演示", "USE_MOCK_FALLBACK=true"),
    ("③ 输入校验（publish.js validate）", "申报人/位置/楼栋/房间/故障类型/手机号正则/描述 7 项校验，缺项 Toast 提示", "/[1][3,4,5,7,8][0-9]{9}$/"),
    ("④ 订阅消息双向通知", "提交报修→管理员；处理完成→报修人，云函数 applyNotice/handleNotice", "wx.requestSubscribeMessage"),
    ("⑤ 图片上传+预览", "wx.chooseMedia 选图，云存储 uploadFile，wx.previewImage 预览，最多 6 张", "cloudPath: repair/xxx.jpg"),
    ("⑥ 异常处理全覆盖（实验七新增）", "所有云调用 try/catch + Toast 反馈，详情不存在/加载失败/权限拒绝均有提示", "catch → wx.showToast"),
]
for i, (t, d, code) in enumerate(snippets):
    row, col = i // 2, i % 2
    x, y = Inches(0.7 + col * 6.1), Inches(2.1 + row * 1.65)
    add_rect(s, x, y, Inches(5.9), Inches(1.5), LIGHT_BG)
    add_rect(s, x, y, Emu(50800), Inches(1.5), PRIMARY if i % 2 == 0 else ACCENT)
    add_text(s, x + Inches(0.15), y + Inches(0.05), Inches(5.6), Inches(0.4),
             t, size=15, color=DARK, bold=True)
    add_text(s, x + Inches(0.15), y + Inches(0.45), Inches(5.6), Inches(0.65),
             d, size=12, color=GRAY, line_spacing=1.3)
    add_text(s, x + Inches(0.15), y + Inches(1.1), Inches(5.6), Inches(0.35),
             "▌ " + code, size=11, color=ACCENT, font="Consolas")

# =========================================================
# 第 7 页：测试与边界
# =========================================================
s = add_slide(); fill_shape(s, WHITE)
add_header(s, "六、测试与边界处理", 7)

add_text(s, Inches(0.7), Inches(1.4), Inches(12), Inches(0.5),
        "边界与异常用例", size=20, color=PRIMARY, bold=True)

# 表格
cases = [
    ("场景", "预期", "实测"),
    ("列表为空", "显示空状态", "mock 3 条，切「已处理」见 1 条"),
    ("无云环境", "自动本地兜底不报错", "关闭云开发可验证"),
    ("提交缺申报人", "Toast 拦截", "validate() 返回 false"),
    ("手机号格式错", "Toast 提示", "正则校验拦截"),
    ("图片超 6 张", "Toast 拦截", "remaining ≤ 0"),
    ("加载失败", "Toast 反馈", "断网/云函数异常"),
    ("详情不存在", "Toast 提示", "删除后回退详情"),
    ("非管理员点管理", "Toast「暂无权限」", "isAdmin 判断"),
]
tbl_x, tbl_y = Inches(0.7), Inches(2.1)
col_w = [Inches(3.2), Inches(3.5), Inches(5.25)]
row_h = Inches(0.5)
for ri, row in enumerate(cases):
    cx = tbl_x
    for ci, txt in enumerate(row):
        is_head = (ri == 0)
        bg = PRIMARY if is_head else (LIGHT_BG if ri % 2 == 1 else WHITE)
        add_rect(s, cx, tbl_y + ri * row_h, col_w[ci], row_h, bg)
        add_text(s, cx + Inches(0.1), tbl_y + ri * row_h, col_w[ci] - Inches(0.2), row_h,
                 txt, size=13, color=WHITE if is_head else DARK,
                 bold=is_head, anchor=MSO_ANCHOR.MIDDLE)
        cx += col_w[ci]

add_text(s, Inches(0.7), Inches(6.7), Inches(12), Inches(0.5),
         "→ 所有边界均有明确 UI 反馈，不静默失败；云失败自动回退本地，主流程不中断",
         size=14, color=ACCENT, bold=True)

# =========================================================
# 第 8 页：项目演示
# =========================================================
s = add_slide(); fill_shape(s, WHITE)
add_header(s, "七、项目演示路径", 8)

add_text(s, Inches(0.7), Inches(1.4), Inches(12), Inches(0.5),
         "5 步演示完整业务闭环（3-5 分钟）", size=20, color=PRIMARY, bold=True)

steps = [
    ("1", "首页筛选", "切「未处理/已处理」Tab，切楼栋筛选，触底分页加载"),
    ("2", "提交报修", "填表（位置/楼栋/房间/手机/故障类型/描述）+ 选图，提交回首页见新记录"),
    ("3", "查看详情", "进详情页，预览图片、一键拨号、复制电话"),
    ("4", "管理后台", "处理一条（状态变「已处理」）、删除一条，验证权限判断"),
    ("5", "无云演示", "关闭云开发重启，验证本地兜底 + mock 种子 + 我的报修缓存正常"),
]
for i, (n, t, d) in enumerate(steps):
    y = Inches(2.2 + i * 0.95)
    # 序号圆
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), y, Inches(0.6), Inches(0.6))
    circ.fill.solid(); circ.fill.fore_color.rgb = PRIMARY
    circ.line.fill.background(); circ.shadow.inherit = False
    add_text(s, Inches(0.8), y, Inches(0.6), Inches(0.6), n,
             size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(1.6), y, Inches(3), Inches(0.6), t,
             size=18, color=DARK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(4.8), y, Inches(8), Inches(0.6), d,
             size=15, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)

add_rect(s, Inches(0.7), Inches(6.8), Inches(11.95), Inches(0.5), LIGHT_BG)
add_text(s, Inches(0.7), Inches(6.8), Inches(11.95), Inches(0.5),
         "演示顺序对应「数据从哪里来→状态如何变→事件如何流转→错误如何处理」四问",
         size=13, color=GRAY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# =========================================================
# 第 9 页：答辩四问
# =========================================================
s = add_slide(); fill_shape(s, WHITE)
add_header(s, "八、答辩四问自检", 9)

add_text(s, Inches(0.7), Inches(1.4), Inches(12), Inches(0.5),
         "任务书要求：能回答四问", size=20, color=PRIMARY, bold=True)

qa = [
    ("数据在哪里？", "云数据库 c_apply（报修记录）+ c_role（角色）+ c_share（分享）\n本地 wx.storage：my_repairs（我的报修）+ mock_seeded（种子）+ admin（管理员）", PRIMARY),
    ("状态在哪里？", "页面级：Page 的 data + setData 驱动渲染\n全局级：App 实例 + wx.storage 持久化\n状态字段：未处理 / 已处理", ACCENT),
    ("事件从哪里来、到哪里去？", "用户操作（点击/输入）→ Page 方法 → utils/api.js（Repository）\n→ 云开发 / 本地兜底 → 回调 setData → UI 刷新", RGBColor(0x4C, 0xAF, 0x50)),
    ("错误在哪里处理？", "api.js 每个方法 try/catch，云失败自动回退本地\n页面层 Toast 明确反馈（加载失败/无权限/校验拦截）\n不静默失败，主流程不中断", RGBColor(0x55, 0x7A, 0xB0)),
]
for i, (q, a, c) in enumerate(qa):
    row, col = i // 2, i % 2
    x, y = Inches(0.7 + col * 6.1), Inches(2.1 + row * 2.5)
    add_rect(s, x, y, Inches(5.9), Inches(2.3), LIGHT_BG)
    add_rect(s, x, y, Inches(5.9), Inches(0.5), c)
    add_text(s, x + Inches(0.15), y, Inches(5.6), Inches(0.5), q,
             size=18, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x + Inches(0.15), y + Inches(0.6), Inches(5.6), Inches(1.6), a,
             size=13, color=DARK, line_spacing=1.4)

# =========================================================
# 第 10 页：总结与致谢
# =========================================================
s = add_slide(); fill_shape(s, PRIMARY)
for cx, cy, d in [(11.5, 0.8, 2.2), (12.0, 5.5, 1.6), (0.3, 6.0, 1.4)]:
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(cy), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0x7B, 0x5F, 0xB0)
    c.line.fill.background(); c.shadow.inherit = False

add_text(s, Inches(1.0), Inches(1.5), Inches(11), Inches(0.8),
         "总结", size=40, color=WHITE, bold=True)
add_rect(s, Inches(1.0), Inches(2.3), Inches(1.5), Emu(38100), ACCENT)

add_bullet(s, Inches(1.0), Inches(2.6), Inches(11.3), [
    "完成 4 页面微信小程序，覆盖任务书 8 项基本功能",
    "云开发 + 本地兜底双模，无云环境也能演示",
    "6 项核心代码：Repository/双模/校验/订阅/上传/异常",
    "8 个边界用例全覆盖，明确 UI 反馈不静默失败",
    "任务书「项目九」8 项基本功能 100% 达成",
], size=18, color=WHITE)

add_rect(s, Inches(1.0), Inches(5.8), Inches(11.3), Inches(1.1), RGBColor(0x7B, 0x5F, 0xB0))
add_text(s, Inches(1.0), Inches(5.85), Inches(11.3), Inches(0.4),
         "答辩四问已能独立解释本人负责模块、核心代码与关键设计决策",
         size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(1.0), Inches(6.3), Inches(11.3), Inches(0.5),
         "感谢老师与同学指导 · 感谢开源项目 MrYe443 提供 UI 参考",
         size=14, color=RGBColor(0xE0, 0xCB, 0xF0), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# 保存
out = r"d:\rain_android\docs\实验七_校园报修小程序_答辩.pptx"
prs.save(out)
print("PPT 已生成:", out)
print("幻灯片数:", len(prs.slides._sldIdLst))
